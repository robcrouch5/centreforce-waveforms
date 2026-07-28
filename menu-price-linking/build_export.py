#!/usr/bin/env python3
"""Add a live 'Photoshop' export tab to the Soultown master, and write a
ready-to-run menu-data.csv.

The Photoshop tab has one row per priced item with a formula that pulls the
price straight from the Menu sheet, so it updates automatically. You then
File > Save As > CSV (with the Photoshop tab active) to feed update-menus.jsx.
"""
import csv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from parse_master import parse, SRC

OUT_XLSX = "Soultown_Menu_2026_linked.xlsx"
OUT_CSV = "menu-data.csv"
ARIAL = "Arial"
PURPLE = "5B2A86"
PINK = "E5197E"

recs = parse()

# ---- price formula: whole pounds show as £8, otherwise £7.90 ---------------
def price_formula(coord):
    ref = f"Menu!{coord}"
    return (f'=IF({ref}="","",IF({ref}=INT({ref}),'
            f'"£"&TEXT({ref},"0"),"£"&TEXT({ref},"0.00")))')

# ---- format the same price for the immediate CSV (Python side) -------------
def price_text(v):
    if isinstance(v, (int, float)):
        return f"£{int(v)}" if float(v) == int(v) else f"£{v:.2f}"
    return str(v)

# =========================== build the .xlsx tab ===========================
wb = load_workbook(SRC)
if "Photoshop" in wb.sheetnames:
    del wb["Photoshop"]
ws = wb.create_sheet("Photoshop")

ws["A1"] = "Photoshop export — do not retype prices here; they pull from the Menu tab automatically."
ws["A1"].font = Font(name=ARIAL, size=12, bold=True, color=PINK)
ws["A2"] = ('When prices change: edit the Menu tab as normal, then with THIS tab open do '
            'File > Save As > CSV UTF-8 and overwrite menu-data.csv, then run update-menus.jsx in Photoshop.')
ws["A2"].font = Font(name=ARIAL, size=10, italic=True, color="808080")
ws["A3"] = ('"Key" = the name to give the matching Photoshop text layer. "File" routes the row to a PSD '
            'whose filename contains that tag (normal / vip). Same key on the horizontal & vertical board updates both.')
ws["A3"].font = Font(name=ARIAL, size=10, italic=True, color="808080")

headers = ["File", "Key (Photoshop layer name)", "Product", "Size", "Price"]
hrow = 5
for c, h in enumerate(headers, start=1):
    cell = ws.cell(hrow, c, h)
    cell.font = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=PURPLE)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

r = hrow + 1
for rec in recs:
    ws.cell(r, 1, rec["file"]).font = Font(name=ARIAL, size=11)
    ws.cell(r, 2, rec["key"]).font = Font(name=ARIAL, size=11, color="333333")
    ws.cell(r, 3, rec["product"]).font = Font(name=ARIAL, size=11)
    ws.cell(r, 4, rec["size"]).font = Font(name=ARIAL, size=11)
    pc = ws.cell(r, 5, price_formula(rec["coord"]))
    pc.font = Font(name=ARIAL, size=11, bold=True)
    pc.alignment = Alignment(horizontal="center")
    for c in range(1, 6):
        ws.cell(r, c).border = border
    r += 1

for i, w in enumerate([10, 44, 34, 10, 10], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = ws.cell(hrow + 1, 1)
wb.save(OUT_XLSX)

# =========================== build the .csv ================================
with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(headers)
    for rec in recs:
        w.writerow([rec["file"], rec["key"], rec["product"], rec["size"], price_text(rec["value"])])

print(f"Wrote {OUT_XLSX} ('Photoshop' tab, {len(recs)} rows) and {OUT_CSV}")
