#!/usr/bin/env python3
"""Parse the Soultown master (Menu sheet) into per-menu price lists.

The sheet is laid out as one column-block per bar. Within a block, size labels
(Pint/Half/25ml/... ) act as headers that apply to the rows beneath them until
the next label. We track the current size per column and emit one record per
priced cell:  (file, key, product, size, coord, value).
"""
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SRC = "Soultown_Menu_2026.xlsx"

# Which master bar feeds which board. file-tag must appear in the PSD filename.
BLOCKS = {
    "normal": (1, 7),    # Bar 1 - Vol 1  (cols A..G)
    "vip":    (32, 38),  # Bar 7 - VIP    (cols AF..AL)
}

SIZE_RE = re.compile(r"(ml\b|\bpint\b|\bhalf\b|\bbottle\b|\bcan\b|\bglass\b|£|\d\s*-\s*\d)", re.I)

def is_size_header(v):
    return isinstance(v, str) and SIZE_RE.search(v) is not None

def slug(s):
    s = str(s).lower().strip()
    s = s.replace("&", "and")
    s = re.sub(r"[^a-z0-9]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    # collapse consecutive duplicate words, e.g. "..._can_can" -> "..._can"
    parts, out = s.split("_"), []
    for w in parts:
        if not out or out[-1] != w:
            out.append(w)
    return "_".join(out)

def parse():
    wb = load_workbook(SRC, data_only=True)
    ws = wb["Menu"]
    out = []
    for file_tag, (c0, c1) in BLOCKS.items():
        # locate category / product / abv columns from the header row (row 2)
        cat_col = prod_col = None
        abv_cols = set()
        for c in range(c0, c1 + 1):
            h = ws.cell(2, c).value
            hs = str(h).strip().lower() if h else ""
            if hs == "category": cat_col = c
            elif hs == "product": prod_col = c
            elif hs.startswith("abv"): abv_cols.add(c)
        price_cols = [c for c in range(c0, c1 + 1)
                      if c not in (cat_col, prod_col) and c not in abv_cols]

        current_size = {c: None for c in price_cols}
        cur_product = None
        for r in range(2, ws.max_row + 1):
            p = ws.cell(r, prod_col).value if prod_col else None
            if p is not None and str(p).strip() != "":
                cur_product = str(p).strip()
            for c in price_cols:
                v = ws.cell(r, c).value
                if v is None or str(v).strip() == "":
                    continue
                if is_size_header(v):
                    current_size[c] = str(v).strip()
                    continue
                # ABV lives in an unlabelled column on some bars; no real price is
                # under £1, so treat sub-£1 numbers as ABV and skip them.
                if isinstance(v, (int, float)) and v < 1:
                    continue
                if isinstance(v, (int, float)) and cur_product:
                    size = current_size[c] or ""
                    body = slug(cur_product + ("_" + size if size else ""))
                    key = f"{file_tag}_{body}"
                    out.append({
                        "file": file_tag, "key": key, "product": cur_product,
                        "size": size, "coord": f"{get_column_letter(c)}{r}",
                        "value": v,
                    })
    return out

if __name__ == "__main__":
    recs = parse()
    for tag in ("normal", "vip"):
        rows = [r for r in recs if r["file"] == tag]
        print(f"\n===== {tag.upper()}  ({len(rows)} priced cells) =====")
        for r in rows:
            print(f"  {r['key']:<42} {r['size']:<8} £{r['value']:<7} <- {r['coord']}  ({r['product']})")
